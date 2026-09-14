import openpyxl
import json
import re
import os

# ===== 解析 AB.xlsx 数据 =====
wb = openpyxl.load_workbook('AB.xlsx', data_only=True)
records = []

for sheet_name in wb.sheetnames:
    year_match = re.search(r'(\d{4})', sheet_name)
    if not year_match:
        continue
    year = int(year_match.group(1))
    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]

    month_cols = {}
    for i, h in enumerate(headers):
        if h is None or i == 0 or i == 1:
            continue
        m = re.match(r'(\d+)月(应交|已交)', str(h))
        if m:
            month = int(m.group(1))
            typ = m.group(2)
            month_cols[i] = (month, typ)
        m2 = re.match(r'(\d+)-(\d+)月(应交|已交)', str(h))
        if m2:
            month = int(m2.group(2))
            typ = m2.group(3)
            month_cols[i] = (month, typ)

    for row in ws.iter_rows(min_row=2, values_only=True):
        room = row[0]
        if room is None:
            continue
        room = str(room).strip()
        if not room:
            continue

        monthly = {}
        for col_i, (month, typ) in month_cols.items():
            val = row[col_i]
            if val is not None:
                try:
                    val = float(val)
                except:
                    val = None
            if month not in monthly:
                monthly[month] = {'应交': None, '已交': None}
            if val is not None:
                monthly[month][typ] = round(val, 2)

        for month, data in monthly.items():
            if data['应交'] is not None or data['已交'] is not None:
                date_str = f'{year}-{month:02d}'
                records.append({
                    'd': date_str,
                    'r': room,
                    'p': data['应交'],
                    'a': data['已交']
                })

rooms = sorted(set(r['r'] for r in records))
data_json = json.dumps(records, ensure_ascii=False, separators=(',', ':'))
rooms_json = json.dumps(rooms, ensure_ascii=False)

print(f'Records: {len(records)}, Rooms: {len(rooms)}')

# ===== 生成 HTML =====
html_template = open('electricity_template.html', 'r', encoding='utf-8').read()
html_out = html_template.replace('__DATA_JSON__', data_json).replace('__ROOMS_JSON__', rooms_json)

with open('electricity_query.html', 'w', encoding='utf-8') as f:
    f.write(html_out)

size = os.path.getsize('electricity_query.html')
print(f'HTML written: {size} bytes = {size/1024:.1f} KB')
print('Done!')
